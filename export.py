# -*- coding: utf-8 -*-
"""导出 Excel：全部房源 / 预算内 / 小区汇总 / 数据说明"""
import os
import sqlite3
from datetime import datetime

import pandas as pd

DIR = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(DIR, "fang.db")
OUT = os.path.join(DIR, "嘉兴二手房房源数据.xlsx")

COLS = {
    "house_id": "房源ID", "city": "城市", "district2": "区县", "district": "区县(列表页)",
    "block": "板块", "community": "小区", "title": "标题",
    "layout": "户型", "rooms": "室", "halls": "厅", "baths": "卫",
    "area": "建筑面积(㎡)", "orientation": "朝向", "floor_desc": "楼层",
    "total_floor": "总楼层", "total_price": "总价(万)", "unit_price": "单价(元/㎡)",
    "build_year": "建成年代", "decoration": "装修", "elevator": "电梯",
    "elevator_ratio": "梯户比", "purpose": "房屋用途", "heating": "供暖方式",
    "structure": "建筑结构", "mortgage": "抵押信息", "list_date": "挂牌时间",
    "last_trade": "上次交易", "url": "房源链接", "source": "数据来源",
}


def main():
    con = sqlite3.connect(DB)
    df = pd.read_sql("SELECT * FROM listings", con)
    con.close()

    for c in COLS:
        if c not in df.columns:
            df[c] = ""
    df["区县"] = df["district2"].fillna("").where(df["district2"].fillna("") != "", df["district"])
    df["区县"] = df["区县"].replace("", "未标注")
    df["房龄"] = df["build_year"].apply(
        lambda x: (datetime.now().year - int(float(x))) if str(x).strip().replace(".", "").isdigit() else None)

    out = pd.DataFrame({
        "序号": range(1, len(df) + 1),
        "房源ID": df["house_id"], "城市": df["city"], "区县": df["区县"], "板块": df["block"],
        "小区": df["community"], "标题": df["title"], "户型": df["layout"],
        "室": df["rooms"], "厅": df["halls"], "卫": df["baths"],
        "建筑面积(㎡)": pd.to_numeric(df["area"], errors="coerce"),
        "朝向": df["orientation"], "楼层": df["floor_desc"], "总楼层": df["total_floor"],
        "总价(万)": pd.to_numeric(df["total_price"], errors="coerce"),
        "单价(元/㎡)": pd.to_numeric(df["unit_price"], errors="coerce"),
        "建成年代": df["build_year"], "房龄(年)": df["房龄"], "装修": df["decoration"],
        "电梯": df["elevator"], "梯户比": df["elevator_ratio"], "房屋用途": df["purpose"],
        "供暖方式": df["heating"], "建筑结构": df["structure"], "抵押信息": df["mortgage"],
        "挂牌时间": df["list_date"], "上次交易": df["last_trade"],
        "房源链接": df["url"], "数据来源": df["source"],
    })

    # 派生指标
    comm_med = out.groupby("小区")["单价(元/㎡)"].transform("median")
    out["小区中位单价"] = comm_med.round(0)
    out["单价偏离小区中位"] = (out["单价(元/㎡)"] / comm_med - 1).round(4)
    out["同小区在售数"] = out.groupby("小区")["房源ID"].transform("count")
    out["疑似重复挂牌数"] = out.groupby(["小区", "建筑面积(㎡)", "总价(万)"])["房源ID"].transform("count")
    out = out.sort_values(["区县", "小区", "总价(万)"], ascending=[True, True, True], na_position="last")

    budget = out[out["总价(万)"] <= 45].sort_values("总价(万)")

    # 去重：同小区+同面积+同总价视为同一套房源的多中介重复挂牌
    dedup = out.drop_duplicates(subset=["小区", "建筑面积(㎡)", "总价(万)"], keep="first")
    dedup = dedup.sort_values(["总价(万)"], ascending=True)

    g = out.dropna(subset=["小区"]).groupby("小区").agg(
        房源数=("房源ID", "count"),
        最低总价万=("总价(万)", "min"),
        中位总价万=("总价(万)", "median"),
        中位单价=("单价(元/㎡)", "median"),
        最小面积=("建筑面积(㎡)", "min"),
        最大面积=("建筑面积(㎡)", "max"),
        区县=("区县", "first"), 板块=("板块", "first"),
    ).reset_index().sort_values("房源数", ascending=False)

    note = pd.DataFrame({"项目": [
        "数据来源", "采集时间", "采集方式", "记录总数", "去重后记录数", "预算内(≤45万)记录数",
        "覆盖范围", "字段说明", "已知局限", "免责声明", "配套工具"],
        "说明": [
            "房天下（搜房网）嘉兴二手房频道 jx.esf.fang.com",
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            "结构化数据接口采集（页面内嵌 schema.org JSON-LD），非人工录入",
            f"{len(out)} 条（含同一房源多中介重复挂牌）",
            f"{len(dedup)} 条（按 小区+面积+总价 去重）",
            f"{len(budget)} 条",
            "嘉兴全市：南湖、秀洲、经济开发区、嘉善、桐乡、平湖、海宁、海盐及周边板块",
            "区县/板块/小区/户型/室厅卫/面积/朝向/楼层/总楼层/总价/单价/建成年代/房龄/装修/电梯/梯户比/房屋用途/供暖/结构/抵押/挂牌时间/链接，另含小区中位单价、单价偏离度、同小区在售数、疑似重复挂牌数",
            "1) 平台标注全站 45232 套，实际因多中介重复挂牌，去重后为「去重后」表数量；"
            "2) 建成年代房天下仅约 3.5% 房源标注，缺失较多，建议向中介或小区物业核实；"
            "3) 楼层、装修、梯户比等来自详情页，部分房源平台本身未填；"
            "4) 挂牌价不等于成交价，实际可议价；"
            "5) 数据为采集时点快照，房源可能已下架或成交",
            "数据仅供个人购房决策参考，不得用于商业用途或对外提供；房源真实性以产权核验和实地看房为准",
            "crawl_fang.py / expand.py / enrich.py / refill.py / export.py（本地脚本，可复跑更新）",
        ],
    })

    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        out.to_excel(w, sheet_name="全部房源", index=False)
        dedup.to_excel(w, sheet_name="去重后房源", index=False)
        budget.to_excel(w, sheet_name="预算内(45万以下)", index=False)
        g.to_excel(w, sheet_name="小区汇总", index=False)
        note.to_excel(w, sheet_name="数据说明", index=False)

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        for ws in w.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F5597")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for i, col in enumerate(ws.columns, 1):
                width = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 42)
            if ws.title != "数据说明":
                ws.auto_filter.ref = ws.dimensions

    print(f"[export] {len(out)} 条 -> {OUT}")
    print(f"[export] 预算内 {len(budget)} 条；小区 {len(g)} 个")


if __name__ == "__main__":
    main()
